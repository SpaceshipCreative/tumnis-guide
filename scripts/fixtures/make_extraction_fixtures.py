"""Generates backend/fixtures/extraction/* (P1-16). Run with a scratch venv holding
reportlab, python-docx, openpyxl, python-pptx, pillow. Deterministic where the libraries
allow (fixed creation dates)."""

import csv
import io
import random
import sys
from datetime import datetime
from pathlib import Path

from docx import Document as Docx
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image as RLImage,
    PageBreak,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
STYLES = getSampleStyleSheet()
H1, H2, BODY = STYLES["Heading1"], STYLES["Heading2"], STYLES["BodyText"]


def _invariant(doc):
    doc.creator = "Tumnis fixtures"
    doc.author = "Tumnis fixtures"
    return doc


def _pdf(name: str, story: list, **kw) -> None:
    doc = SimpleDocTemplate(
        str(OUT / name), pagesize=A4, title=name, author="Tumnis fixtures", invariant=1, **kw
    )
    doc.build(story)


LOREM = (
    "Tumnis keeps the plan, the brief and the files of a project in one place, so every "
    "agent works from the same facts. This paragraph is filler text written for the "
    "extraction fixtures; it carries no meaning beyond giving the page a body of prose."
)


def text_2p() -> None:
    _pdf(
        "text-2p.pdf",
        [
            Paragraph("Project charter", H1),
            Paragraph("Introduction", H2),
            Paragraph(
                "The Acme site redesign starts in March and replaces the marketing pages "
                "with a faster, accessible build.",
                BODY,
            ),
            Paragraph(LOREM, BODY),
            PageBreak(),
            Paragraph("Scope", H2),
            Paragraph(
                "The scope covers the home page, the pricing page and the blog; the shop "
                "stays out of scope until the second phase.",
                BODY,
            ),
            Paragraph(LOREM, BODY),
        ],
    )


def rate_card() -> None:
    table = Table(
        [
            ["Role", "Hourly rate (EUR)"],
            ["Senior designer", "160"],
            ["Designer", "120"],
            ["Developer", "140"],
            ["Project manager", "110"],
        ],
        colWidths=[8 * cm, 5 * cm],
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.8, colors.black),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ]
        )
    )
    _pdf(
        "rate-card-table.pdf",
        [
            Paragraph("Acme rate card", H1),
            Paragraph("Overview", H2),
            Paragraph(
                "This rate card lists what Acme pays for design and build work in 2026. "
                "Rates are per hour and exclude VAT.",
                BODY,
            ),
            Paragraph(LOREM, BODY),
            PageBreak(),
            Paragraph("Rates", H2),
            Paragraph("Hourly rates by role:", BODY),
            Spacer(1, 0.4 * cm),
            table,
            PageBreak(),
            Paragraph("Terms", H2),
            Paragraph(
                "Invoices are due within thirty days. Travel is billed at cost.", BODY
            ),
            Paragraph(LOREM, BODY),
        ],
    )


def two_column() -> None:
    path = OUT / "two-column.pdf"
    doc = BaseDocTemplate(
        str(path), pagesize=A4, title="two-column.pdf", author="Tumnis fixtures", invariant=1
    )
    width, height = A4
    margin, gap = 2 * cm, 1 * cm
    col = (width - 2 * margin - gap) / 2
    frames = [
        Frame(margin, margin, col, height - 2 * margin, id="left"),
        Frame(margin + col + gap, margin, col, height - 2 * margin, id="right"),
    ]
    doc.addPageTemplates([PageTemplate(id="two", frames=frames)])
    from reportlab.platypus import FrameBreak

    doc.build(
        [
            Paragraph("Garden notes", H2),
            Paragraph(
                "The left column is about the community garden: tomatoes went in on "
                "Saturday and the compost bins were turned.",
                BODY,
            ),
            Paragraph(LOREM, BODY),
            FrameBreak(),
            Paragraph("Bicycle club", H2),
            Paragraph(
                "The right column is about the bicycle club: the Sunday ride leaves the "
                "harbour at nine and returns by noon.",
                BODY,
            ),
            Paragraph(LOREM, BODY),
        ]
    )


def _text_image(lines: list[str], size=(1240, 1754), font_size=36) -> Image.Image:
    img = Image.new("L", size, 255)
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, font_size)
    y = 120
    for line in lines:
        draw.text((110, y), line, fill=0, font=font)
        y += int(font_size * 1.6)
    return img


def scanned() -> None:
    img = _text_image(
        [
            "Delivery note",
            "",
            "Twelve chairs were delivered to the",
            "Acme office on Monday morning.",
            "The driver asked for a signature.",
        ]
    )
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    _pdf("scanned-1p.pdf", [RLImage(buf, width=A4[0] - 4 * cm, height=(A4[0] - 4 * cm) * 1.414)])


def handwriting() -> None:
    rnd = random.Random(16)
    img = Image.new("L", (1240, 1754), 255)
    draw = ImageDraw.Draw(img)
    for row in range(14):
        y = 160 + row * 100
        x = 100
        while x < 1100:
            points = [(x + i * 6, y + rnd.randint(-18, 18)) for i in range(rnd.randint(4, 12))]
            draw.line(points, fill=0, width=3)
            x = points[-1][0] + rnd.randint(18, 40)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    _pdf(
        "handwriting.pdf",
        [
            Paragraph("Meeting notes", H1),
            Paragraph("The typed summary is on this page; the scan follows.", BODY),
            PageBreak(),
            RLImage(buf, width=A4[0] - 4 * cm, height=(A4[0] - 4 * cm) * 1.414),
        ],
    )


def brief_docx() -> None:
    doc = Docx()
    _invariant(doc.core_properties)
    doc.core_properties.created = datetime(2026, 3, 1)
    doc.core_properties.modified = datetime(2026, 3, 1)
    doc.add_heading("Client brief", level=1)
    doc.add_heading("Goals", level=2)
    doc.add_paragraph("Double the newsletter sign-ups by June without paid ads.")
    doc.add_heading("Audience", level=2)
    doc.add_paragraph("Small studios that run their own bookings and invoices.")
    doc.save(OUT / "brief.docx")


def budget_xlsx() -> None:
    wb = Workbook()
    wb.properties.creator = "Tumnis fixtures"
    wb.properties.created = datetime(2026, 3, 1)
    wb.properties.modified = datetime(2026, 3, 1)
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Item", "Amount"])
    summary.append(["Design", 4800])
    summary.append(["Build", 7200])
    summary.append(["Hosting", 600])
    details = wb.create_sheet("Details")
    details.append(["Task", "Hours"])
    details.append(["Wireframes", 12])
    details.append(["Visual design", 18])
    wb.save(OUT / "budget.xlsx")


def kickoff_pptx() -> None:
    prs = Presentation()
    prs.core_properties.author = "Tumnis fixtures"
    prs.core_properties.created = datetime(2026, 3, 1)
    prs.core_properties.modified = datetime(2026, 3, 1)
    layout = prs.slide_layouts[1]
    for title, body in (
        ("Kickoff", "Welcome to the Acme redesign kickoff."),
        ("Timeline", "Design in March, build in April, launch in May."),
        ("Next steps", "Share the brand guide and book the content workshop."),
    ):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(OUT / "kickoff.pptx")


def notes_md() -> None:
    (OUT / "notes.md").write_text(
        "# Workshop notes\n\n"
        "## Decisions\n\n"
        "- The launch date moves to the first Monday in May.\n"
        "- The blog keeps its current URLs.\n\n"
        "## Open questions\n\n"
        "Who writes the pricing page copy?\n"
    )


def rates_csv() -> None:
    with (OUT / "rates.csv").open("w", newline="") as out:
        writer = csv.writer(out)
        writer.writerow(["Role", "Hourly rate"])
        writer.writerow(["Senior designer", "160"])
        writer.writerow(["Copywriter", "95"])


def receipt_png() -> None:
    img = _text_image(
        ["Corner Cafe", "", "2 x Flat white   7.00", "1 x Croissant    3.50", "", "Total 10.50"],
        size=(900, 700),
        font_size=40,
    )
    img.save(OUT / "receipt.png", format="PNG", optimize=True)


def html_named_pdf() -> None:
    (OUT / "html-named.pdf").write_text(
        "<!DOCTYPE html>\n<html>\n<head><title>Invoice</title></head>\n"
        "<body><h1>Invoice</h1><p>This is an HTML page saved with a .pdf name.</p>"
        "<script>alert(1)</script></body>\n</html>\n"
    )


for make in (
    text_2p,
    rate_card,
    two_column,
    scanned,
    handwriting,
    brief_docx,
    budget_xlsx,
    kickoff_pptx,
    notes_md,
    rates_csv,
    receipt_png,
    html_named_pdf,
):
    make()
print("ok")
